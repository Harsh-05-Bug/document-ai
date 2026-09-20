"""Turn an uploaded file into a list of {page_number, text} records.

Page numbers are what make citations useful, so they are preserved
wherever the format has a real notion of a page. Formats that don't
(DOCX, TXT, CSV) report None and the UI just cites the document.
"""
import csv
import logging

from pypdf import PdfReader

log = logging.getLogger(__name__)


def extract_pdf(path: str) -> list[dict]:
    reader = PdfReader(path)
    pages = []
    for index, page in enumerate(reader.pages):
        text = page.extract_text() or ""
        if not text.strip():
            text = _ocr_page(path, index)
        pages.append({"page_number": index + 1, "text": text})
    return pages


def _ocr_page(path: str, index: int) -> str:
    """Scanned pages have no text layer. OCR them if the extras are installed."""
    try:
        import pytesseract
        from pdf2image import convert_from_path
    except ImportError:
        log.info("page %s has no text layer and OCR extras are not installed", index + 1)
        return ""

    try:
        images = convert_from_path(path, first_page=index + 1, last_page=index + 1, dpi=200)
        return pytesseract.image_to_string(images[0]) if images else ""
    except Exception:
        log.exception("OCR failed on page %s", index + 1)
        return ""


def extract_docx(path: str) -> list[dict]:
    import docx

    document = docx.Document(path)
    parts = [p.text for p in document.paragraphs if p.text.strip()]

    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))

    return [{"page_number": None, "text": "\n".join(parts)}]


def extract_xlsx(path: str) -> list[dict]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    pages = []
    for sheet_index, sheet in enumerate(workbook.worksheets, start=1):
        lines = []
        for row in sheet.iter_rows(values_only=True):
            cells = [str(c) for c in row if c is not None]
            if cells:
                lines.append(" | ".join(cells))
        if lines:
            # One "page" per sheet so citations can name the sheet.
            pages.append({"page_number": sheet_index, "text": f"Sheet: {sheet.title}\n" + "\n".join(lines)})
    workbook.close()
    return pages


def extract_csv(path: str) -> list[dict]:
    with open(path, newline="", encoding="utf-8", errors="ignore") as handle:
        rows = list(csv.reader(handle))
    if not rows:
        return []
    header, *body = rows
    lines = [" | ".join(header)] + [" | ".join(row) for row in body]
    return [{"page_number": None, "text": "\n".join(lines)}]


def extract_txt(path: str) -> list[dict]:
    with open(path, encoding="utf-8", errors="ignore") as handle:
        return [{"page_number": None, "text": handle.read()}]


EXTRACTORS = {
    "application/pdf": extract_pdf,
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": extract_docx,
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": extract_xlsx,
    "text/csv": extract_csv,
}


def extract(path: str, mime_type: str) -> list[dict]:
    extractor = EXTRACTORS.get(mime_type, extract_txt)
    pages = extractor(path)
    return [p for p in pages if p["text"].strip()]


def suffix_for(mime_type: str) -> str:
    return {
        "application/pdf": ".pdf",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document": ".docx",
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": ".xlsx",
        "text/csv": ".csv",
    }.get(mime_type, ".txt")
